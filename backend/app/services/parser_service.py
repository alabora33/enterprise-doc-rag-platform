from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


ParsedSection = dict[str, Any]


def parse_txt(file_path: str) -> list[ParsedSection]:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        text = file.read()

    return [
        {
            "text": text,
            "page_number": None,
            "sheet_name": None,
            "source_type": "txt",
        }
    ]


def parse_pdf(file_path: str) -> list[ParsedSection]:
    reader = PdfReader(file_path)
    sections: list[ParsedSection] = []

    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            sections.append(
                {
                    "text": text,
                    "page_number": page_index,
                    "sheet_name": None,
                    "source_type": "pdf",
                }
            )

    return sections


def parse_docx(file_path: str) -> list[ParsedSection]:
    document = DocxDocument(file_path)

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text and paragraph.text.strip()
    ]

    text = "\n".join(paragraphs)

    return [
        {
            "text": text,
            "page_number": None,
            "sheet_name": None,
            "source_type": "docx",
        }
    ]


def parse_xlsx(file_path: str) -> list[ParsedSection]:
    workbook = load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    sections: list[ParsedSection] = []

    for sheet in workbook.worksheets:
        rows_as_text: list[str] = []

        for row in sheet.iter_rows(values_only=True):
            values = [
                str(cell).strip()
                for cell in row
                if cell is not None and str(cell).strip()
            ]

            if values:
                rows_as_text.append(" | ".join(values))

        sheet_text = "\n".join(rows_as_text)

        if sheet_text.strip():
            sections.append(
                {
                    "text": sheet_text,
                    "page_number": None,
                    "sheet_name": sheet.title,
                    "source_type": "xlsx",
                }
            )

    return sections


def parse_document(file_path: str) -> list[ParsedSection]:
    extension = Path(file_path).suffix.lower()

    if extension == ".txt":
        return parse_txt(file_path)

    if extension == ".pdf":
        return parse_pdf(file_path)

    if extension == ".docx":
        return parse_docx(file_path)

    if extension == ".xlsx":
        return parse_xlsx(file_path)

    raise ValueError(f"Unsupported file extension for parsing: {extension}")